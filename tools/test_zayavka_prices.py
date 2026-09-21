"""Заявка магазину с ценами (владелец, 22 сен 2026): после названия — наша
закупочная цена, справа от Total — сумма строки, внизу две строки — TOTAL
(количества по районам, итог растянут на Total+Amount) и AMOUNT (сумма
каждого района = чек магазина, и общий итог).

  • колонки: №, Code, Item, Price, пять районов, Total, Amount;
  • цена — из stock_value.cost_map (та же, что оценивает склад); нет цены —
    клетка пустая, сумма строки пустая, в итог не входит;
  • сумма строки = цена × Total; сумма района = Σ цена × количество района;
    общий итог = Σ сумм строк = Σ районов; всё формулами и с посчитанным
    значением рядом (для превью телефона);
  • ответ магазина читается как раньше: те же количества по районам, строки
    TOTAL и AMOUNT не превращаются в позиции.
"""
import asyncio, io, os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("MONGO_URI", ""); os.environ.setdefault("AMBAR_OWNER_IDS", "1")
import logging
logging.disable(logging.CRITICAL)
from mongomock_motor import AsyncMongoMockClient
from openpyxl import load_workbook
import db, supply_routes as sr, stock_value

FAIL = []
def eq(name, got, want):
    ok = got == want
    print(("  ok  " if ok else "  FAIL") + f" {name}: {got!r}" + ("" if ok else f" ≠ {want!r}"))
    if not ok: FAIL.append(name)

D = ["jvc", "bbay", "silicon", "alguses", "tecom"]
ROWS = [  # id, name, количества по районам
    ("p1", "Absolut 1 ltr", [10, 21, 16, 15, 14]),
    ("p31", "Heineken 0.33 can", [5, 7, 0, 4, 3]),
    ("p62", "Malfy Rosa 0.7 ltr", [1, 1, 1, 0, 0]),          # цены нет
    ("p21", "Blue Label 1 ltr", [0, 1, 0, 0, 0]),
]
COSTS = {"p1": 34.65, "p31": 79.8, "p21": 850}


async def fake_rows(day):
    rows = [{"id": i, "name": n, "need_total": sum(q),
             "cells": {o: {"need": v} for o, v in zip(D, q)}} for i, n, q in ROWS]
    return {"day": "2026-09-22", "districts": [{"id": o} for o in D], "rows": rows, "all_rows": rows}


async def fake_costs():
    return dict(COSTS)


def unwrap(h):
    while hasattr(h, "__wrapped__"):
        h = h.__wrapped__
    return h


class _Field:
    name = "file"
    def __init__(self, raw): self.raw = raw
    async def read(self, decode=False): return self.raw


class _Reader:
    def __init__(self, raw): self.left = [_Field(raw)]
    async def next(self): return self.left.pop() if self.left else None


class _Req(dict):
    def __init__(self, raw): super().__init__(owner_id=1); self.raw = raw
    async def multipart(self): return _Reader(self.raw)


async def main():
    db._db = AsyncMongoMockClient()["ambar_zayavka"]
    sr._order_rows = fake_rows
    stock_value.cost_map = fake_costs
    raw, name = await sr._build_book("")
    eq("имя файла", name, "AMBAR-zayavka-2026-09-22.xlsx")
    wf = load_workbook(io.BytesIO(raw))["Order"]
    wv = load_workbook(io.BytesIO(raw), data_only=True)["Order"]
    hdr = [wv.cell(row=3, column=c).value for c in range(2, 13)]
    eq("колонки", hdr, ["№", "Code", "Item", "Price, AED", "B1 JVC", "B2 Business Bay", "B3 Silicon Oasis",
                        "B4 Al Qusais", "B5 Tecom", "Total", "Amount, AED"])
    got = {wv.cell(row=r, column=3).value: [wv.cell(row=r, column=c).value for c in (5, 11, 12)] for r in range(4, 8)}
    eq("Absolut: цена, Total, сумма", got["p1"], [34.65, 76, 2633.4])
    eq("Heineken: коробка × 19", got["p31"], [79.8, 19, 1516.2])
    eq("без цены — пустая цена и пустая сумма", (got["p62"][0], got["p62"][1], got["p62"][2] in (None, "")),
       (None, 3, True))
    eq("сумма строки — формулой от цены и Total", wf["L4"].value, '=IF(OR(E4="",N(K4)=0),"",E4*K4)')
    # итоговые строки
    tot = [wv.cell(row=8, column=c).value for c in range(2, 12)]
    eq("TOTAL: количества по районам и итог", (tot[0], tot[4:10]), ("TOTAL", [16, 30, 17, 19, 17, 99]))
    amt = [wv.cell(row=9, column=c).value for c in range(2, 12)]
    want = [round(sum(COSTS.get(i, 0) * q[k] for i, _, q in ROWS), 2) for k in range(5)]
    eq("AMOUNT: сумма каждого района = цена × количество района", (amt[0], amt[4:9]), ("AMOUNT, AED", want))
    eq("общий итог = сумма строк = сумма районов", (amt[9], round(2633.4 + 1516.2 + 850, 2), round(sum(want), 2)),
       (4999.6, 4999.6, 4999.6))
    merged = {str(r) for r in wf.merged_cells.ranges}
    eq("итоги растянуты на Total+Amount", {"K8:L8", "K9:L9"} <= merged, True)
    eq("сумма района — формулой", wf["F9"].value, "=SUMPRODUCT($E$4:$E$7,F4:F7)")
    eq("общий итог — формулой", wf["K9"].value, "=SUM(L4:L7)")

    print("── ответ магазина читается как раньше ─────────────────────────")
    resp = await unwrap(sr.handle_import)(_Req(raw))
    import json
    body = json.loads(resp.text)
    eq("импорт прошёл", (resp.status, body.get("ok")), (200, True))
    sup = await db.supply_get(body["supply_id"])
    bd = {it["id"]: {o: it["by_district"].get(o, 0) for o in D} for it in sup["items"]}
    eq("позиции — только товар, без TOTAL/AMOUNT", sorted(bd), sorted(i for i, _, _ in ROWS))
    eq("количества по районам те же", {i: [bd[i][o] for o in D] for i in bd}, {i: q for i, _, q in ROWS})
    eq("чужих кодов нет", sup.get("unknown"), [])

    print("ИТОГ:", "все прошли" if not FAIL else f"провалено {len(FAIL)}: {FAIL}")
    return 1 if FAIL else 0


sys.exit(asyncio.run(main()))
