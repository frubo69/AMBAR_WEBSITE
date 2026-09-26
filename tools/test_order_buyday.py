"""Заявка датируется днём закупки, а не уходящей сменой (владелец, 26 сен 2026).

Смена идёт с 10:00 до 10:00, а закрывают её под утро следующего дня —
замерено по журналу: файл уходил в магазин между 05:50 и 09:17. Датировался
он при этом днём СМЕНЫ, и утром 26-го владелец видел документ с числом 25:
«возникает ощущение, что передо мной старая заявка». Везут по ней в тот же
день, когда она ушла, — этим днём её и датируем.

  • день закупки = день смены + 1, и считается от дня смены, а не по стенным
    часам: закрыли смену вечером того же дня — дата не должна уехать назад;
  • он же стоит в заголовке листа, в имени файла и в подписи к документу;
  • у живой (ещё не замороженной) заявки он тоже есть — экран не должен
    показывать сменой одно, а файлом другое;
  • у замороженной берётся записанный, а не пересчитывается;
  • подпись к файлу вырезается из имени целиком (была отрезана первая цифра
    года: «Заявка в магазин · 026-09-25»).

    python3 tools/test_order_buyday.py
"""
import asyncio, os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ["MONGO_URI"] = ""; os.environ.setdefault("AMBAR_OWNER_IDS", "1")
import logging; logging.disable(logging.WARNING)
from mongomock_motor import AsyncMongoMockClient                   # noqa: E402
import db, stock_routes as sr, supply_routes as sup                # noqa: E402

FAIL = []
def eq(имя, дали, ждём):
    ок = дали == ждём
    print(("  ok  " if ок else "  FAIL") + f" {имя}: {дали!r}" + ("" if ок else f" ≠ {ждём!r}"))
    if not ок: FAIL.append(имя)

СМЕНА, ЗАКУПКА, РАЙОН = "2026-09-26", "2026-09-27", "silicon"


async def подменить(have):
    def пусто():
        return {"have": {}, "have_exact": {}, "sug": {}, "came": {}, "gone": {}, "counted": True}
    async def _base(day):
        out = {o: пусто() for o in sr.OFFICE_IDS}
        out[РАЙОН] = dict(пусто(), have=dict(have), have_exact=dict(have),
                          sug={p: 10 for p in have})
        return out
    sr._district_base = _base


async def main():
    db._db = AsyncMongoMockClient()["ambar_order_buyday"]
    sr._biz_day = lambda *a, **k: СМЕНА
    pid = next(iter(sr._catalog()))
    await подменить({pid: 2})

    print("── живая заявка, снимка ещё нет ───────────────────────────────")
    sr.base_drop()
    d = await sr.order_rows(СМЕНА)
    eq("день остался днём смены", d["day"], СМЕНА)
    eq("ДЕНЬ ЗАКУПКИ — СЛЕДУЮЩИЙ", d["buy_day"], ЗАКУПКА)

    print("── файл в магазин ─────────────────────────────────────────────")
    raw, name = await sup._build_book(СМЕНА)
    eq("имя файла по дню закупки", name, f"AMBAR-zayavka-{ЗАКУПКА}.xlsx")
    eq("дня смены в имени нет", СМЕНА in name, False)
    from openpyxl import load_workbook
    import io
    ws = load_workbook(io.BytesIO(raw))["Order"]
    eq("заголовок листа", ws["B1"].value, f"Purchase order · {ЗАКУПКА}")

    print("── подпись к документу ────────────────────────────────────────")
    подпись = name.removeprefix("AMBAR-zayavka-").removesuffix(".xlsx")
    eq("вырезана дата целиком", подпись, ЗАКУПКА)

    print("── заморозка ──────────────────────────────────────────────────")
    r = await sr.freeze_order(СМЕНА)
    eq("заморожена", r["ok"], True)
    eq("ЗАПИСАН ДЕНЬ ЗАКУПКИ", r["buy_day"], ЗАКУПКА)
    снимок = await db.zayavka_freeze_get(СМЕНА)
    eq("и он лежит в снимке", снимок.get("buy_day"), ЗАКУПКА)

    print("── после заморозки берём записанный, а не считаем заново ──────")
    await db._db.zayavki.update_one({"_id": СМЕНА}, {"$set": {"buy_day": "2026-01-01"}})
    sr.base_drop()
    d = await sr.order_rows(СМЕНА)
    eq("уважаем записанный день", d["buy_day"], "2026-01-01")
    raw2, name2 = await sup._build_book(СМЕНА)
    eq("и файл идёт за ним", name2, "AMBAR-zayavka-2026-01-01.xlsx")

    print("── чужой день: стенные часы тут не при чём ────────────────────")
    # Смена мая, а на дворе сентябрь. Пока день закупки брался по стенным
    # часам, здесь оказывалось сегодняшнее число — и заявка за давнюю смену
    # датировалась сегодняшним днём.
    ДАВНО, ДАВНО_ЗАКУПКА = "2026-05-10", "2026-05-11"
    r = await sr.freeze_order(ДАВНО)
    eq("день закупки от дня смены, а не от часов", r["buy_day"], ДАВНО_ЗАКУПКА)
    sr.base_drop()
    d = await sr.order_rows(ДАВНО)
    eq("и в ответе он же", d["buy_day"], ДАВНО_ЗАКУПКА)

    print("── конец месяца и года ────────────────────────────────────────")
    from bizday import next_day
    eq("30 сентября → 1 октября", next_day("2026-09-30"), "2026-10-01")
    eq("31 декабря → 1 января", next_day("2026-12-31"), "2027-01-01")

    print(("\nПРОВАЛЫ: " + ", ".join(FAIL)) if FAIL else "\nвсё сошлось")
    return 1 if FAIL else 0


sys.exit(asyncio.run(main()))
