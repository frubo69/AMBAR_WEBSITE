"""Цены магазина из файла заявки и документ заявки (владелец, 24 окт 2026).

  • снимок «что просили» (zayavka_save) не стирает заморозку дня (base, buy_day);
  • из файла магазина читаются цены «Price, AED» — по коду, названию, номеру;
  • разница с нашей закупочной — предложение: принять → закупочная = цена
    магазина (рука владельца), оставить → закрыто; одинаковая цена — не разница.

    python3 tools/test_price_proposals.py
"""
import asyncio, io, json, os, sys
from datetime import datetime, timezone
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ["MONGO_URI"] = ""; os.environ.setdefault("AMBAR_OWNER_IDS", "1")
import logging; logging.disable(logging.CRITICAL)
from mongomock_motor import AsyncMongoMockClient                  # noqa: E402
from aiohttp.test_utils import make_mocked_request                # noqa: E402
import db, supply_routes as sr, stock_value as sv                 # noqa: E402

FAIL = []
def eq(name, got, want):
    ok = got == want
    print(("  ok  " if ok else "  FAIL") + f" {name}: {got!r}" + ("" if ok else f" ≠ {want!r}"))
    if not ok: FAIL.append(name)

async def main():
    db._db = AsyncMongoMockClient()["price_props"]

    print("── снимок заявки не стирает заморозку ─────────────────────────")
    ok = await db.zayavka_freeze_set("2026-10-23", {"p1": {"bbay": {"calc": 5}}}, "2026-10-24", {"bbay": {"counted": True}})
    await db.zayavka_save("2026-10-23", {"p1": 5}, {"p1": {"bbay": 5}})
    fr = await db.zayavka_freeze_get("2026-10-23")
    eq("заморозка цела после снимка", (ok, bool(fr.get("base")), fr.get("buy_day")), (True, True, "2026-10-24"))
    eq("снимок на месте", (await db.zayavka_last_full()).get("asked"), {"p1": 5})
    eq("ищется по дню закупки", (await db.zayavka_freeze_for_buy("2026-10-24")).get("day"), "2026-10-23")

    print("── цены из файла ──────────────────────────────────────────────")
    cat = sr._catalog_by_id()
    pids = list(cat)[:3]
    names = [cat[p]["name"] for p in pids]
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = sr.SHEET_MAIN
    ws.append(["Логотип магазина"])
    ws.append([sr.NUM_COL, sr.ITEM_COL, sr.PRICE_COL, "B2 Business Bay", sr.TOTAL_COL])
    ws.append([1, names[0], 65, 3, 3])              # цена изменена
    ws.append([2, names[1], "62,00", 2, 2])         # та же цена, текстом с запятой
    ws.append([3, names[2].upper() + " ", None, 1, 1])  # цены нет
    ws.append(["TOTAL", "", "", 6, 6])
    raw = io.BytesIO(); wb.save(raw)
    from openpyxl import load_workbook
    ws2 = load_workbook(io.BytesIO(raw.getvalue()), data_only=True).worksheets[0]
    hdr, cols = sr._header_row(ws2)
    eq("шапка найдена ниже логотипа", hdr, 2)
    prices = sr._file_prices(ws2, hdr, cols, cat)
    eq("цены по позициям: изменённая и текстовая, без пустой", prices, {pids[0]: 65.0, pids[1]: 62.0})

    async def _costs(): return {pids[0]: 62.0, pids[1]: 62.0}
    sv.cost_map = _costs
    changes = await sv.price_changes(prices)
    eq("разница только там, где цена другая", [(c["id"], c["ours"], c["theirs"]) for c in changes], [(pids[0], 62.0, 65.0)])
    await sv.price_proposals_add(changes, "S1")
    pend = await sv.price_proposals_pending()
    eq("предложение ждёт", [(r["id"], r["state"], r["sid"]) for r in pend], [(pids[0], "pending", "S1")])
    r = await sv.price_proposal_decide(pids[0], True, "Владелец")
    eq("принято: закупочная стала ценой магазина (рукой)", (r["state"], (await db.cost_overrides()).get(pids[0], {}).get("price")), ("accepted", 65.0))
    eq("нерешённых не осталось", await sv.price_proposals_pending(), [])
    eq("второе решение — None", await sv.price_proposal_decide(pids[0], False, "В"), None)
    await sv.price_proposals_add([{"id": pids[1], "name": names[1], "ours": 62.0, "theirs": 70.0, "unit": 1}], "S2")
    r = await sv.price_proposal_decide(pids[1], False, "Владелец")
    eq("оставили: цена не тронута", (r["state"], pids[1] in await db.cost_overrides()), ("kept", False))

    print("── ручка решения ──────────────────────────────────────────────")
    await sv.price_proposals_add([{"id": pids[2], "name": names[2], "ours": None, "theirs": 40.0, "unit": 1}], "S3")
    h = sv.handle_proposal
    while hasattr(h, "__wrapped__"): h = h.__wrapped__
    rq = make_mocked_request("POST", "/x"); rq["owner_id"] = 1
    rq._read_bytes = json.dumps({"all": True, "accept": True, "as": "Владелец"}).encode()
    resp = await h(rq); j = json.loads(resp.text)
    eq("принять все: пусто и цена записана", (resp.status, j["proposals"], (await db.cost_overrides()).get(pids[2], {}).get("price")), (200, [], 40.0))

    print("ИТОГ:", "все прошли" if not FAIL else f"провалено {len(FAIL)}: {FAIL}")
    return 1 if FAIL else 0

sys.exit(asyncio.run(main()))
